from datetime import date
from io import BytesIO
import unittest
import uuid
from unittest.mock import patch

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api.applications import export_applications
from app.core.database import Base
from app.models.job_application import ApplicationStatus, WorkMode
from app.models.user import User
from app.repositories import job_application as repo
from app.services import export_service, intent_dispatcher


class TestExportService(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(bind=self.engine)
        self.Session = sessionmaker(bind=self.engine)
        self.session = self.Session()

        self.user_a_id = uuid.uuid4()
        self.user_b_id = uuid.uuid4()
        self.session.add(User(id=self.user_a_id, line_user_id="line_user_a"))
        self.session.add(User(id=self.user_b_id, line_user_id="line_user_b"))
        self.session.commit()

    async def asyncTearDown(self):
        self.session.close()
        Base.metadata.drop_all(bind=self.engine)
        self.engine.dispose()

    def test_export_empty_list(self):
        stream = export_service.export_applications_to_xlsx([])
        self.assertIsInstance(stream, BytesIO)

        # Verify workbook structure
        wb = load_workbook(stream)
        ws = wb.active
        self.assertEqual(ws.title, "Job Applications")
        self.assertEqual(ws.max_row, 1)  # Only header
        self.assertEqual(ws.cell(row=1, column=1).value, "วันที่สมัคร")
        self.assertEqual(ws.cell(row=1, column=2).value, "บริษัท")

    def test_export_with_applications(self):
        app1 = repo.create_application(self.session, self.user_a_id, {
            "company": "KBank",
            "position": "Backend Developer",
            "status": ApplicationStatus.APPLIED,
            "work_mode": WorkMode.HYBRID,
            "salary": "40,000",
            "date_applied": date(2026, 10, 1),
            "note": "รอบแรกผ่านแล้ว",
        })
        app2 = repo.create_application(self.session, self.user_a_id, {
            "company": "SCB",
            "position": "Data Scientist",
            "status": ApplicationStatus.INTERVIEWED,
            "work_mode": WorkMode.REMOTE,
            "salary": "55,000",
            "date_applied": date(2026, 9, 25),
        })

        apps = [app1, app2]
        stream = export_service.export_applications_to_xlsx(apps)

        wb = load_workbook(stream)
        ws = wb.active
        self.assertEqual(ws.max_row, 3)  # Header + 2 rows

        # Check row 1 (KBank)
        self.assertEqual(ws.cell(row=2, column=2).value, "KBank")
        self.assertEqual(ws.cell(row=2, column=3).value, "Backend Developer")
        self.assertEqual(ws.cell(row=2, column=4).value, "สมัครแล้ว")
        self.assertEqual(ws.cell(row=2, column=5).value, "Hybrid")
        self.assertEqual(ws.cell(row=2, column=6).value, "40,000")
        self.assertEqual(ws.cell(row=2, column=10).value, "รอบแรกผ่านแล้ว")

        # Check row 2 (SCB)
        self.assertEqual(ws.cell(row=3, column=2).value, "SCB")
        self.assertEqual(ws.cell(row=3, column=4).value, "สัมภาษณ์แล้ว")

    def test_export_endpoint_response(self):
        repo.create_application(self.session, self.user_a_id, {
            "company": "Agoda",
            "position": "Fullstack Engineer",
            "status": ApplicationStatus.APPLIED,
        })

        response = export_applications(
            status=None,
            company=None,
            position=None,
            work_mode=None,
            date_from=None,
            date_to=None,
            db=self.session,
        )

        self.assertEqual(
            response.media_type,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            "attachment; filename=job_applications.xlsx",
            response.headers["Content-Disposition"],
        )

    @patch("app.services.gemini_service.parse_intent_with_gemini")
    async def test_export_intent_dispatcher(self, mock_parse):
        repo.create_application(self.session, self.user_a_id, {
            "company": "LINE Man",
            "position": "Product Manager",
            "status": ApplicationStatus.ACCEPTED,
        })

        mock_parse.return_value = {
            "type": "function_call",
            "name": "export_applications",
            "args": {},
        }

        response_text = await intent_dispatcher.dispatch_user_message(
            db=self.session,
            user_id=self.user_a_id,
            user_message="ขอ export ข้อมูลการสมัครงานเป็น excel หน่อยครับ",
        )

        self.assertIn("พบข้อมูลการสมัครงานทั้งหมด 1 รายการ", response_text)
        self.assertIn("/applications/export", response_text)


if __name__ == "__main__":
    unittest.main()
