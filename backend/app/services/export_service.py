from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.job_application import JobApplication


def export_applications_to_xlsx(applications: list[JobApplication]) -> BytesIO:
    """Generate an Excel (.xlsx) file in-memory containing the given job applications."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Job Applications"

    # Define headers
    headers = [
        "วันที่สมัคร",
        "บริษัท",
        "ตำแหน่ง",
        "สถานะ",
        "รูปแบบงาน",
        "เงินเดือน",
        "สถานที่ทำงาน",
        "แหล่งที่มา",
        "ลิงก์ประกาศงาน",
        "บันทึกเพิ่มเติม",
        "วันที่บันทึกระบบ",
    ]

    ws.append(headers)

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    ws.row_dimensions[1].height = 28

    # Data styling
    data_font = Font(name="Calibri", size=10)
    data_alignment_left = Alignment(horizontal="left", vertical="center")
    data_alignment_center = Alignment(horizontal="center", vertical="center")

    # Add application rows
    for row_idx, app in enumerate(applications, start=2):
        status_text = app.status.value if hasattr(app.status, "value") else str(app.status)
        work_mode_text = app.work_mode.value if app.work_mode and hasattr(app.work_mode, "value") else (str(app.work_mode) if app.work_mode else "-")
        date_applied_text = str(app.date_applied) if app.date_applied else "-"
        created_at_text = app.created_at.strftime("%Y-%m-%d %H:%M") if app.created_at else "-"

        row_data = [
            date_applied_text,
            app.company or "-",
            app.position or "-",
            status_text,
            work_mode_text,
            app.salary or "-",
            app.location or "-",
            app.source or "-",
            app.job_url or "-",
            app.note or "-",
            created_at_text,
        ]

        ws.append(row_data)

        # Style data cells
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            # Center dates and status, left-align text
            if col_idx in [1, 4, 5, 11]:
                cell.alignment = data_alignment_center
            else:
                cell.alignment = data_alignment_left

        ws.row_dimensions[row_idx].height = 22

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_length = 0
        for cell in col:
            val_str = str(cell.value or "")
            # Estimate width (give Thai characters a bit more weight)
            val_len = len(val_str.encode("utf-8")) // 2 if any(ord(c) > 127 for c in val_str) else len(val_str)
            if val_len > max_length:
                max_length = val_len
        ws.column_dimensions[col_letter].width = max(max_length + 5, 14)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
